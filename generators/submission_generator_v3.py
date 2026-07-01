from docx import Document


def generate_submission_sheet(
    candidate_data,
    recruiter_data,
    skill_match_data,
    output_path
):

    doc = Document(
        "templates/tcs_submission_template.docx"
    )

    table = doc.tables[0]

    # ==========================
    # Candidate Information
    # ==========================

    table.cell(1, 1).text = candidate_data.get(
        "candidate_name",
        ""
    )

    table.cell(2, 1).text = candidate_data.get(
        "phone",
        ""
    )

    table.cell(3, 1).text = candidate_data.get(
        "email",
        ""
    )

    table.cell(4, 1).text = candidate_data.get(
        "linkedin",
        ""
    )

    table.cell(5, 1).text = candidate_data.get(
        "location",
        ""
    )

    table.cell(6, 1).text = recruiter_data.get(
        "willing_to_relocate",
        ""
    )

    table.cell(7, 1).text = recruiter_data.get(
        "former_tcs_employee",
        ""
    )

    # Leave Row 8 untouched
    # It contains TCS instructions

    table.cell(9, 1).text = recruiter_data.get(
        "interview_availability",
        ""
    )

    table.cell(10, 1).text = recruiter_data.get(
        "general_availability",
        ""
    )

    # ==========================
    # Submission Skills
    # ==========================

    submission_skills = skill_match_data.get(
        "submission_skills",
        []
    )

    for index, skill in enumerate(
        submission_skills[:3]
    ):

        row_number = 13 + index

        table.cell(
            row_number,
            0
        ).text = skill.get(
            "skill",
            ""
        )

        table.cell(
            row_number,
            1
        ).text = skill.get(
            "years_experience",
            ""
        )

        table.cell(
            row_number,
            2
        ).text = skill.get(
            "relevant_experience",
            ""
        )

    doc.save(
        output_path
    )